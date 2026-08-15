import json
import math
from decimal import Decimal, InvalidOperation
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse
from django.views.decorators.http import require_POST
from django.core.paginator import Paginator
from .models import (Mahsulot, Kategoriya, SkannerLog, MahsulotSharh, SharhRasm,
                     BarkodKatalog, Inventarizatsiya, InventarizatsiyaItem, NarxTarix)
from .services import ReviewService, ReviewStats
from users.models import DokonEgasiProfil
from users.services import ServiceError, json_ok, json_err


def _decimal(value, default='0'):
    try:
        return Decimal(str(value if value not in (None, '') else default))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("Narx noto'g'ri kiritilgan")


# Kasrli (tarozi) birliklar — kg/litr kabi
KASRLI_BIRLIKLAR = ('kg', 'litr', 'gramm', 'metr')


def _miqdor(value, birlik='dona', default='0'):
    """Miqdorni Decimal qilib qaytaradi. dona → butun, kg/litr → 3 xona. Manfiy=0."""
    from decimal import ROUND_HALF_UP
    try:
        d = Decimal(str(value if value not in (None, '') else default))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("Miqdor noto'g'ri kiritilgan")
    if d < 0:
        d = Decimal('0')
    if (birlik or '').lower() in KASRLI_BIRLIKLAR:
        return d.quantize(Decimal('0.001'))
    return d.quantize(Decimal('1'), rounding=ROUND_HALF_UP)


def _son(d):
    """Miqdorni chiroyli matn: butun→'3', kasrli→'0.75'."""
    try:
        d = Decimal(str(d))
    except (InvalidOperation, ValueError, TypeError):
        return str(d)
    if d == d.to_integral_value():
        return format(d.quantize(Decimal('1')), 'f')
    return format(d.normalize(), 'f')


def _get_dokon(user):
    if user.rol == 'dokon_egasi':
        try:
            return user.dokon_profil
        except DokonEgasiProfil.DoesNotExist:
            return None
    elif user.rol == 'sotuvchi':
        try:
            return user.sotuvchi_profil.dokon
        except Exception:
            return None
    return None


def _egasi_dokon(user):
    """Faqat DO'KON EGASI uchun dokon qaytaradi. Sotuvchi/boshqa -> None.
    Narx, inventarizatsiya kabi sezgir amallar uchun — sotuvchi omborni
    o'zgartirib o'g'irlikni yashira olmasligi kerak."""
    if user.rol == 'dokon_egasi':
        try:
            return user.dokon_profil
        except DokonEgasiProfil.DoesNotExist:
            return None
    return None


@login_required
def mahsulot_royxat(request):
    dokon = None
    if request.user.rol == 'dokon_egasi':
        dokon = _get_dokon(request.user)
        mahsulotlar = Mahsulot.objects.filter(dokon=dokon) if dokon else []
    elif request.user.rol == 'sotuvchi':
        dokon = _get_dokon(request.user)
        mahsulotlar = Mahsulot.objects.filter(dokon=dokon) if dokon else []
    elif request.user.rol == 'zavod':
        try:
            mahsulotlar = Mahsulot.objects.filter(zavod=request.user.zavod_profil)
        except Exception:
            mahsulotlar = []
    else:
        mahsulotlar = Mahsulot.objects.filter(faol=True)

    # Aqlli bashorat: do'kon egasi/sotuvchi uchun "taxminan qachon tugaydi"
    if dokon:
        from kassa.hisobot_service import mahsulotlar_bashorati
        mahsulotlar = list(mahsulotlar)
        bashorat_map = mahsulotlar_bashorati(dokon, mahsulotlar)
        for m in mahsulotlar:
            m.bashorat = bashorat_map.get(m.id)

    return render(request, 'mahsulot/royxat.html', {'mahsulotlar': mahsulotlar})


@login_required
def mahsulot_qoshish(request):
    if request.user.rol not in ['dokon_egasi', 'zavod']:
        return redirect('login')

    if request.method == 'POST':
        try:
            nomi = request.POST.get('nomi', '').strip()
            narx = _decimal(request.POST.get('narx', '0'))
            sotib_olish_narxi = _decimal(request.POST.get('sotib_olish_narxi', '0'))
            birlik = request.POST.get('birlik', 'dona')
            soni = _miqdor(request.POST.get('soni', '0'), birlik)
            minimum_son = _miqdor(request.POST.get('minimum_son', '5'), birlik, default='5')
            tavsif = request.POST.get('tavsif', '')
            barkod = request.POST.get('barkod', '').strip() or None
            kategoriya_slug = request.POST.get('kategoriya', '').strip()

            # Yaroqlilik muddati (ixtiyoriy, sana+soat: YYYY-MM-DDTHH:MM; eski YYYY-MM-DD ham qabul qilinadi)
            yaroqlilik = request.POST.get('yaroqlilik_muddati', '').strip() or None
            yaroqlilik_muddati = None
            if yaroqlilik:
                from datetime import datetime as _dt
                from django.utils import timezone as _tz
                for _fmt in ('%Y-%m-%dT%H:%M', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
                    try:
                        _naive = _dt.strptime(yaroqlilik, _fmt)
                        yaroqlilik_muddati = _tz.make_aware(_naive, _tz.get_current_timezone())
                        break
                    except ValueError:
                        continue

            # Mahsulot rasmi majburiy (#13)
            if 'rasm' not in request.FILES:
                return JsonResponse({'status': 'error', 'error': 'Mahsulot rasmi majburiy! Iltimos, rasm yuklang.'}, status=400)

            m = Mahsulot(
                nomi=nomi,
                narx=narx,
                sotib_olish_narxi=sotib_olish_narxi,
                soni=max(0, soni),
                minimum_son=max(0, minimum_son),
                birlik=birlik,
                tavsif=tavsif,
                barkod=barkod,
                yaroqlilik_muddati=yaroqlilik_muddati,
            )

            if kategoriya_slug:
                try:
                    m.kategoriya = Kategoriya.objects.get(nomi=kategoriya_slug)
                except Kategoriya.DoesNotExist:
                    pass

            if request.user.rol == 'dokon_egasi':
                m.dokon = _get_dokon(request.user)
            elif request.user.rol == 'zavod':
                try:
                    m.zavod = request.user.zavod_profil
                except Exception:
                    pass

            if 'rasm' in request.FILES:
                m.rasm = request.FILES['rasm']

            # Shu do'konda bu barkod allaqachon bormi? (do'kon ichida unikal)
            if barkod and m.dokon_id and \
                    Mahsulot.objects.filter(dokon=m.dokon, barkod=barkod).exists():
                return JsonResponse({
                    'status': 'error',
                    'error': "Bu barkodli mahsulot do'koningizda allaqachon bor.",
                }, status=400)

            m.save()
            # Umumiy barkod katalogini boyitamiz (boshqa do'konlarga foyda)
            if barkod:
                BarkodKatalog.yangila(barkod, nomi, kategoriya_slug)
            return JsonResponse({'status': 'ok', 'id': m.id})

        except ValueError as e:
            # Foydalanuvchiga tushunarli validatsiya xabari
            return JsonResponse({'status': 'error', 'error': str(e)}, status=400)
        except Exception:
            # Texnik xatolik — foydalanuvchiga stack trace ko'rsatmaymiz (#12)
            return JsonResponse({'status': 'error', 'error': "Nimadir xato yuz berdi. Iltimos, ma'lumotlarni tekshirib qayta urinib ko'ring."}, status=400)

    kategoriyalar = Kategoriya.objects.all()
    return render(request, 'mahsulot/qoshish.html', {'kategoriyalar': kategoriyalar})


@login_required
def mahsulot_import(request):
    """Excel (.xlsx) yoki CSV fayldan mahsulotlarni bir yo'la qo'shish.
    Ustunlar: nomi, narx, soni, barkod, kategoriya, birlik (birinchi qator — sarlavha)."""
    if request.user.rol not in ['dokon_egasi', 'zavod']:
        return redirect('login')

    if request.method == 'POST' and request.FILES.get('fayl'):
        fayl = request.FILES['fayl']
        rows = []
        try:
            if fayl.name.lower().endswith('.csv'):
                import csv as _csv
                import io as _io
                text = _io.TextIOWrapper(fayl.file, encoding='utf-8-sig')
                rows = [list(r) for r in _csv.reader(text)]
            else:
                from openpyxl import load_workbook
                wb = load_workbook(fayl, read_only=True, data_only=True)
                rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
        except Exception:
            messages.error(request, "Faylni o'qib bo'lmadi. Excel (.xlsx) yoki CSV yuklang.")
            return redirect('mahsulot_import')

        if len(rows) < 2:
            messages.error(request, "Faylda sarlavha va kamida 1 qator mahsulot bo'lishi kerak.")
            return redirect('mahsulot_import')

        header = [str(c or '').strip().lower() for c in rows[0]]

        def idx(*names):
            for nm in names:
                if nm in header:
                    return header.index(nm)
            return -1

        i_nom = idx('nomi', 'nom', 'name', 'mahsulot', 'название')
        i_narx = idx('narx', 'narxi', 'price', 'sotish narxi', 'цена')
        i_soni = idx('soni', 'son', 'qty', 'miqdor', 'количество')
        i_bar = idx('barkod', 'barcode', 'shtrix', 'штрих')
        i_kat = idx('kategoriya', 'category', 'turi', 'категория')
        i_birlik = idx('birlik', 'unit', 'olchov')

        if i_nom < 0 or i_narx < 0:
            messages.error(request, "Faylda kamida 'nomi' va 'narx' ustunlari bo'lishi shart.")
            return redirect('mahsulot_import')

        dokon = _get_dokon(request.user) if request.user.rol == 'dokon_egasi' else None
        zavod = getattr(request.user, 'zavod_profil', None) if request.user.rol == 'zavod' else None
        qoshildi = otkazildi = 0

        def _hujayra(r, i):
            return r[i] if 0 <= i < len(r) else None

        for r in rows[1:]:
            try:
                nomi = str(_hujayra(r, i_nom) or '').strip()
                if not nomi:
                    otkazildi += 1
                    continue
                narx_raw = _hujayra(r, i_narx)
                narx = _decimal(narx_raw) if narx_raw not in (None, '') else 0
                birlik = str(_hujayra(r, i_birlik) or 'dona').strip() or 'dona'
                soni_raw = _hujayra(r, i_soni)
                soni = _miqdor(soni_raw, birlik) if soni_raw not in (None, '') else Decimal('0')
                barkod = None
                bar_raw = _hujayra(r, i_bar)
                if bar_raw not in (None, ''):
                    barkod = str(bar_raw).strip()
                    if barkod.endswith('.0'):
                        barkod = barkod[:-2]   # Excel raqamni float qiladi
                kat_matn = str(_hujayra(r, i_kat) or '').strip()

                # Do'kon ichida barkod takrorlanmasin
                if barkod and dokon and Mahsulot.objects.filter(dokon=dokon, barkod=barkod).exists():
                    otkazildi += 1
                    continue

                m = Mahsulot(nomi=nomi, narx=narx, soni=max(0, soni),
                             birlik=birlik, barkod=barkod, dokon=dokon, zavod=zavod)
                # Kategoriya: to'g'ridan-to'g'ri slug yoki nom/matndan aniqlash
                kat_slug = ''
                if kat_matn and Kategoriya.objects.filter(nomi=kat_matn).exists():
                    kat_slug = kat_matn
                else:
                    kat_slug = _kategoriya_matndan(f"{kat_matn} {nomi}")
                if kat_slug:
                    m.kategoriya = Kategoriya.objects.filter(nomi=kat_slug).first()
                m.save()
                if barkod:
                    BarkodKatalog.yangila(barkod, nomi, m.kategoriya.nomi if m.kategoriya_id else '')
                qoshildi += 1
            except Exception:
                otkazildi += 1

        messages.success(request, f"✅ {qoshildi} ta mahsulot qo'shildi." +
                         (f" {otkazildi} ta o'tkazib yuborildi." if otkazildi else ""))
        return redirect('mahsulot_royxat')

    return render(request, 'mahsulot/import.html')


@login_required
def skanner_view(request):
    return render(request, 'mahsulot/skanner.html')


@login_required
def jonatma_qabul_dokon(request):
    """Do'kon jo'natma QR kodini skaner qilib omborga qabul qiladigan sahifa."""
    dokon = _get_dokon(request.user)
    if not dokon:
        return redirect('login')
    return render(request, 'mahsulot/jonatma_qabul.html')


@login_required
@require_POST
def jonatma_qabul_dokon_action(request):
    """QR (barkod) orqali jo'natmadagi HAMMA mahsulotni do'kon omboriga qo'shadi."""
    from django.db import transaction
    from django.utils import timezone
    from datetime import datetime, time as _time
    from yetkazib_beruvchi.models import Jonatma

    dokon = _get_dokon(request.user)
    if not dokon:
        return JsonResponse({'error': "Ruxsat yo'q"}, status=403)
    try:
        barkod = (json.loads(request.body).get('barkod') or '').strip().upper()
    except Exception:
        return JsonResponse({'error': 'JSON xato'}, status=400)
    if not barkod:
        return JsonResponse({'error': 'Barkod kiritilmadi'}, status=400)

    j = (Jonatma.objects.filter(barkod=barkod)
         .prefetch_related('mahsulotlar').first())
    if not j:
        return JsonResponse({'error': 'Jo\'natma topilmadi'}, status=404)
    if j.dokon_qabul:
        return JsonResponse({'error': 'Bu jo\'natma allaqachon qabul qilingan'}, status=400)
    # Jo'natma boshqa do'konga mo'ljallangan bo'lsa — qabul qilib bo'lmaydi
    # (begona do'kon zavod tovarini o'z omboriga o'tkazib olmasin).
    manzil = j._manba_dokon()
    if manzil and manzil.id != dokon.id:
        return JsonResponse(
            {'error': "Bu jo'natma boshqa do'konga mo'ljallangan."}, status=403)

    qoshildi = []
    with transaction.atomic():
        for jm in j.mahsulotlar.all():
            muddat = None
            if jm.tugash_sana:
                muddat = timezone.make_aware(datetime.combine(jm.tugash_sana, _time(23, 59)))
            # Do'konda shu barkod/nom bo'yicha mavjud mahsulotni topamiz
            mah = None
            if jm.barkod:
                mah = Mahsulot.objects.filter(dokon=dokon, barkod=jm.barkod).first()
            if not mah:
                mah = Mahsulot.objects.filter(dokon=dokon, nomi=jm.mahsulot_nomi).first()
            if mah:
                mah.soni += jm.soni
                mah.birlik = jm.birlik or mah.birlik
                if muddat:
                    mah.yaroqlilik_muddati = muddat
                mah.faol = True
                mah.save(update_fields=['soni', 'birlik', 'yaroqlilik_muddati', 'faol'])
                yangi = False
            else:
                mah = Mahsulot.objects.create(
                    dokon=dokon, nomi=jm.mahsulot_nomi,
                    barkod=(jm.barkod or None), soni=jm.soni,
                    narx=jm.narx, sotib_olish_narxi=jm.narx,
                    birlik=(jm.birlik or 'dona'), yaroqlilik_muddati=muddat, faol=True,
                )
                yangi = True
            qoshildi.append({'nomi': mah.nomi, 'soni': jm.soni, 'birlik': mah.birlik, 'yangi': yangi})

        j.dokon_qabul = True
        if j.holat != 'yetkazildi':
            j.holat = 'yetkazildi'
        j.save(update_fields=['dokon_qabul', 'holat'])

    return JsonResponse({
        'status': 'ok',
        'jonatma_id': j.id,
        'soni': len(qoshildi),
        'mahsulotlar': qoshildi,
    })


@login_required
@require_POST
def skanner_action(request):
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'error': 'JSON xato'}, status=400)

    dokon = _get_dokon(request.user)
    if not dokon:
        return JsonResponse({'error': "Ruxsat yo'q"}, status=403)

    barkod = data.get('barkod', '').strip()
    harakat = data.get('harakat', 'kirim')

    if harakat not in ['kirim', 'chiqim']:
        return JsonResponse({'error': "Harakat noto'g'ri"}, status=400)

    # MUHIM: faqat O'Z do'koni mahsuloti. Barkod do'konlar aro takrorlanadi
    # (uniq_dokon_barkod) — dokon filtrsiz begona ombor o'zgarardi yoki bir xil
    # barkod bir nechta bo'lsa MultipleObjectsReturned (500) berardi.
    try:
        mahsulot = Mahsulot.objects.get(barkod=barkod, dokon=dokon)
    except Mahsulot.DoesNotExist:
        return JsonResponse({'error': f'Barkod topilmadi: {barkod}'}, status=404)
    except Mahsulot.MultipleObjectsReturned:
        mahsulot = Mahsulot.objects.filter(barkod=barkod, dokon=dokon).first()

    try:
        soni = _miqdor(data.get('soni', 1), mahsulot.birlik, default='1')
    except ValueError:
        return JsonResponse({'error': "Soni noto'g'ri"}, status=400)
    if soni <= 0:
        return JsonResponse({'error': "Soni 0 dan katta bo'lishi kerak"}, status=400)

    # Ombor o'zgarishi DB tomonda F() bilan — poyga (race) sharoitida ham
    # to'g'ri bo'ladi. Chiqimda faqat qoldiq yetsagina kamayadi.
    from django.db.models import F
    if harakat == 'kirim':
        Mahsulot.objects.filter(id=mahsulot.id).update(soni=F('soni') + soni)
    elif harakat == 'chiqim':
        kamaydi = Mahsulot.objects.filter(
            id=mahsulot.id, soni__gte=soni
        ).update(soni=F('soni') - soni)
        if not kamaydi:
            return JsonResponse({'error': f"Yetarli mahsulot yo'q (qoldi: {mahsulot.soni})"}, status=400)

    SkannerLog.objects.create(
        mahsulot=mahsulot,
        harakat=harakat,
        soni=soni,
        foydalanuvchi=request.user,
    )

    mahsulot.refresh_from_db(fields=['soni'])
    return JsonResponse({
        'status': 'ok',
        'mahsulot': mahsulot.nomi,
        'yangi_son': float(mahsulot.soni),
        'kam_qoldi': mahsulot.kam_qoldi,
    })


# Kategoriya kalitlari — OFF teglari va mahsulot nomidan (uz/ru/en/brend).
# Tartib muhim: birinchi moslik g'olib. Aniqroqlari yuqorida.
_KAT_KALITLAR = [
    ('choy_qahva', ['tea', 'coffee', 'chaj', 'chay', 'choy', 'kofe', 'kahva', 'nescafe',
                    'jacobs', 'lipton', 'greenfield', 'curtis', 'ahmad', ' 3 in 1', '3v1', 'kapuchino']),
    ('makaron', ['pasta', 'noodle', 'macaroni', 'spaghetti', 'makaron', 'vermishel',
                 'lapsha', 'rojki', 'makfa', 'spagetti']),
    ('sut', ['milk', 'dairy', 'yogurt', 'yoghurt', 'cheese', 'kefir', 'smetana', 'tvorog',
             'sut ', 'sir', 'pishloq', 'syr', 'ryazhenka', 'slivki', 'prostokvasha', 'snejok',
             'danone', 'epica', 'chudo', 'actimel', 'yogurt', 'moloko', 'qatiq', 'tvorojok']),
    ('ichimliklar', ['beverage', 'drink', 'soda', 'cola', 'kola', 'water', 'juice', 'lemonade',
                     'energy', 'soft-drink', 'napit', 'voda', 'suv', 'sok', 'nektar', 'nectar',
                     'pepsi', 'fanta', 'sprite', 'cola', 'limonad', 'kvas', 'gazli', 'gazirov',
                     'adrenalin', 'red bull', 'gorilla', 'hell', 'damla', 'hydrolife', 'bonaqua',
                     'santal', 'piko', 'fuse', 'rich', 'dobriy', 'water', 'napitok']),
    ('shirinlik', ['chocolate', 'candy', 'sweet', 'biscuit', 'cookie', 'confection',
                   'wafer', 'shokolad', 'shokolad', 'choco', 'konfet', 'pechenye', 'печенье',
                   'vafli', 'marshmallow', 'marmelad', 'zefir', 'halva', 'holva', 'kozinak',
                   'jvachka', 'gum', 'trident', 'orbit', 'dirol', 'oreo', 'kitkat', 'kit kat',
                   'snickers', ' mars', 'twix', 'bounty', 'nutella', 'raffaello', 'rafaello',
                   'milka', 'roshen', 'konti', 'yaskino', 'jaskino', 'rokki', 'shariklar', 'batonchik']),
    ('gazaklar', ['snack', 'chips', 'chipsy', 'chipsi', 'crisps', 'cracker', 'kraker',
                  'popcorn', 'nuts', 'orehi', 'oreshki', 'suxarik', 'сухарик', 'semechki',
                  'fistashka', 'lays', 'cheetos', 'pringles', 'hrusteam', 'korochki', 'snek',
                  'grizzly', 'flint', 'cheers']),
    ('non', ['bread', 'bakery', 'baguette', 'hleb', 'lepeshka', 'baton', 'buhanka', 'hlebsi', ' non ']),
    ('gosht', ['meat', 'sausage', 'beef', 'chicken', 'kolbasa', 'myaso', 'gosht', 'sosiska',
               'sosisk', 'vetchina', 'govyadina', 'kuritsa', 'buzoq', 'hot dog', 'pelmen']),
    ('baliq', ['fish', 'seafood', 'ryba', 'baliq', 'tuna', 'sardina', 'seld', 'kilka',
               'shprot', 'fortuna', 'gorbusha']),
    ('tuxum', ['egg', 'yaic', 'tuxum']),
    ('konserva', ['canned', 'tinned', 'konserv', 'tushenka', 'pashtet']),
    ('yog', ['oil', 'butter', 'margarine', 'maslo', 'oliva', 'olive', 'podsolnech',
             'rastitel', 'margarin', 'yog\'']),
    ('un_yorma', ['flour', 'cereal', 'rice', 'grain', 'muka', 'krupa', 'ris', 'guruch',
                  'grechka', 'grechnev', 'ovsyanka', 'gerkules', 'bulgur', 'manka', 'perlovka', 'un ']),
    ('ziravorlar', ['spice', 'sauce', 'condiment', 'ketchup', 'kechup', 'mayonnaise', 'mayonez',
                    'specii', 'sous', 'gorchitsa', 'mustard', 'sol ', 'perets', 'priprava',
                    'uksus', 'adjika', 'tomat', 'shashlik', 'шашлик']),
    ('meva_sabzavot', ['fruit', 'vegetable', 'frukt', 'ovoshch', 'meva', 'sabzavot', 'fruit-tella']),
]


def _kategoriya_matndan(matn):
    """Ixtiyoriy matndan (nomi/kategoriya teglari) bizning kategoriya slugini topadi."""
    matn = (matn or '').lower()
    for kalit, sozlar in _KAT_KALITLAR:
        if any(s in matn for s in sozlar):
            return kalit
    return ''


def _off_kategoriya(product):
    """Open Food Facts mahsulotidan bizning kategoriya kalitini topadi."""
    matn = ' '.join([
        ' '.join(product.get('categories_tags') or []),
        product.get('categories') or '',
        product.get('product_name') or '',
        product.get('brands') or '',
    ])
    return _kategoriya_matndan(matn)


@login_required
def mahsulot_barkod_info(request):
    """Barkod bo'yicha mahsulot nomi/kategoriyasini qaytaradi.
    1) Avval o'z bazamiz, 2) keyin Open Food Facts (bepul, ochiq baza)."""
    barkod = (request.GET.get('barkod') or '').strip()
    if not barkod:
        return JsonResponse({'topildi': False, 'xato': "Barkod bo'sh"})

    # 1) Umumiy barkod katalogi (barcha do'konlar birgalikda to'ldirgan)
    kat = BarkodKatalog.objects.filter(barkod=barkod).first()
    if kat:
        return JsonResponse({
            'topildi': True,
            'nomi': kat.nomi,
            'kategoriya': kat.kategoriya,
            'manba': 'katalog',
        })

    # 2) Mavjud mahsulotlardan (eski ma'lumot, katalogga hali tushmaganlar)
    mavjud = Mahsulot.objects.filter(barkod=barkod).select_related('kategoriya').first()
    if mavjud:
        return JsonResponse({
            'topildi': True,
            'nomi': mavjud.nomi,
            'kategoriya': mavjud.kategoriya.nomi if mavjud.kategoriya else '',
            'manba': 'baza',
        })

    # 2) Open Food Facts — dunyo bo'yicha ochiq barkod bazasi
    try:
        import requests
        r = requests.get(
            f'https://world.openfoodfacts.org/api/v2/product/{barkod}.json',
            params={'fields': 'product_name,product_name_en,brands,categories_tags,categories'},
            timeout=6,
            headers={'User-Agent': 'Savdoyim/1.0 (savdo tizimi)'},
        )
        j = r.json()
        if j.get('status') == 1 or j.get('product'):
            p = j.get('product', {}) or {}
            nomi = (p.get('product_name') or p.get('product_name_en')
                    or p.get('brands') or '').strip()
            if nomi:
                return JsonResponse({
                    'topildi': True,
                    'nomi': nomi,
                    'kategoriya': _off_kategoriya(p),
                    'manba': 'openfoodfacts',
                })
    except Exception:
        pass

    return JsonResponse({'topildi': False})


def api_mahsulot_qidirish(request):
    # Ochiq qidiruv — mehmon (guest) ham foydalanadi. Faqat faol/mavjud mahsulotlarni o'qiydi.
    q = request.GET.get('q', '').strip()
    lat = request.GET.get('lat')
    lng = request.GET.get('lng')

    # Faqat do'kon mahsulotlari — zavod katalogi (dokon=None) qidiruvga chiqmaydi.
    mahsulotlar = Mahsulot.objects.filter(
        nomi__icontains=q, soni__gt=0, faol=True, dokon__isnull=False
    ).select_related('dokon') if q else []

    natijalar = []
    for m in mahsulotlar:
        masofa = None
        if lat and lng and m.dokon and m.dokon.latitude:
            try:
                dlat = math.radians(float(lat) - m.dokon.latitude)
                dlng = math.radians(float(lng) - m.dokon.longitude)
                a = (math.sin(dlat / 2) ** 2 +
                     math.cos(math.radians(float(lat))) *
                     math.cos(math.radians(m.dokon.latitude)) *
                     math.sin(dlng / 2) ** 2)
                masofa = round(6371 * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a)), 2)
            except Exception:
                pass

        natijalar.append({
            'id': m.id,
            'nomi': m.nomi,
            'narx': float(m.narx),
            'soni': m.soni,
            'rasm': m.rasm.url if m.rasm else None,
            'dokon_nomi': m.dokon.dokon_nomi if m.dokon else '',
            'dokon_user_id': m.dokon.foydalanuvchi_id if m.dokon else None,
            'lat': m.dokon.latitude if m.dokon else None,
            'lng': m.dokon.longitude if m.dokon else None,
            'manzil': m.dokon.manzil if m.dokon else '',
            'masofa': masofa,
        })

    if lat and lng:
        natijalar.sort(key=lambda x: x['masofa'] if x['masofa'] is not None else 99999)

    return JsonResponse({'natijalar': natijalar})


# ═══════════════════════════════════════════════
#  MAHSULOT SHARHLARI — thin views -> ReviewService / ReviewStats
# ═══════════════════════════════════════════════

def mahsulot_detail(request, mahsulot_id):
    """Ochiq mahsulot sahifasi — reyting, statistika, sharhlar (hamma ko'radi).
    Sharh formasi faqat huquqli foydalanuvchiga ko'rsatiladi."""
    mahsulot = (Mahsulot.objects
                .select_related('dokon', 'kategoriya')
                .filter(id=mahsulot_id).first())
    # Zavod katalogidagi mahsulot (dokon=None) ochiq sahifada ko'rinmaydi —
    # marketplace faqat do'kon mahsulotlari uchun.
    if not mahsulot or not mahsulot.dokon_id:
        from django.http import Http404
        raise Http404("Mahsulot topilmadi")

    sharh_qs = (mahsulot.sharhlar.filter(holat='faol')
                .select_related('foydalanuvchi')
                .prefetch_related('rasmlar'))
    page = Paginator(sharh_qs, 10).get_page(request.GET.get('page'))
    stats = ReviewStats.get(mahsulot)
    taqsimot_list = [(s, stats['taqsimot'][s], stats['foizlar'][s]) for s in (5, 4, 3, 2, 1)]

    can_review = False
    mening_sharhim = None
    if request.user.is_authenticated:
        from buyurtma.models import MijozBuyurtma
        mening_sharhim = mahsulot.sharhlar.filter(foydalanuvchi=request.user).first()
        has_order = MijozBuyurtma.objects.filter(
            mijoz=request.user, holat='olib_ketildi', itemlar__mahsulot=mahsulot
        ).exists()
        can_review = has_order and mening_sharhim is None

    return render(request, 'mahsulot/detail.html', {
        'mahsulot': mahsulot,
        'page_obj': page,
        'sharhlar': page.object_list,
        'stats': stats,
        'taqsimot_list': taqsimot_list,
        'can_review': can_review,
        'mening_sharhim': mening_sharhim,
        'is_admin': request.user.is_authenticated and (request.user.is_staff or request.user.is_superuser),
        'yulduzlar': [5, 4, 3, 2, 1],
        'mehmon': not request.user.is_authenticated,   # guest chrome + login CTA
    })


@login_required
@require_POST
def sharh_qoldirish(request):
    """Yangi sharh. Biznes logika ReviewService'da. 400/403/404/409 kodlari."""
    mahsulot = get_object_or_404(Mahsulot, id=request.POST.get('mahsulot_id'))
    try:
        sharh = ReviewService.create(
            request.user, mahsulot,
            request.POST.get('yulduz'), request.POST.get('izoh', ''),
            request.FILES.getlist('rasmlar'), request,
        )
    except ServiceError as e:
        return json_err(e.message, e.errors, e.code)
    return json_ok("Sharhingiz qo'shildi. Rahmat!", {'sharh_id': sharh.id}, status=201)


@login_required
@require_POST
def sharh_tahrirlash(request, sharh_id):
    """Foydalanuvchi FAQAT o'z sharhini tahrirlaydi (egalik ReviewService'da)."""
    sharh = get_object_or_404(MahsulotSharh, id=sharh_id)
    try:
        ReviewService.update(
            request.user, sharh,
            request.POST.get('yulduz', sharh.yulduz), request.POST.get('izoh'),
            request.FILES.getlist('rasmlar'), request,
        )
    except ServiceError as e:
        return json_err(e.message, e.errors, e.code)
    return json_ok("Sharh yangilandi.", {'sharh_id': sharh.id})


@login_required
@require_POST
def sharh_moderatsiya(request, sharh_id):
    """Faqat administrator yashiradi/tiklaydi (ReviewService tekshiradi)."""
    sharh = get_object_or_404(MahsulotSharh, id=sharh_id)
    amal = request.POST.get('amal')  # 'yashir' | 'tikla'
    try:
        if amal == 'yashir':
            ReviewService.hide(request.user, sharh, request)
        elif amal == 'tikla':
            ReviewService.restore(request.user, sharh, request)
        else:
            return json_err("Noto'g'ri amal. 'yashir' yoki 'tikla' bo'lishi kerak.", status=400)
    except ServiceError as e:
        return json_err(e.message, e.errors, e.code)
    return json_ok("Bajarildi.", {'holat': sharh.holat})


# ══════════════════════════════════════════════════════════════
#  MAHSULOT TAHRIRLASH + NARX TARIXI
# ══════════════════════════════════════════════════════════════
@login_required
def mahsulot_tahrir(request, mahsulot_id):
    """FAQAT DO'KON EGASI mahsulotni tahrirlaydi (narx/tannarx/min/faol).
    Narx yoki tannarx o'zgarsa — NarxTarix ga yoziladi. Sotuvchi kira olmaydi
    (narxni o'zgartirib suiiste'mol qilmasligi uchun)."""
    dokon = _egasi_dokon(request.user)
    if not dokon:
        messages.error(request, "Bu amal faqat do'kon egasiga ruxsat etilgan.")
        return redirect('mahsulot_royxat')
    mahsulot = get_object_or_404(Mahsulot, id=mahsulot_id, dokon=dokon)

    if request.method == 'POST':
        from django.db import transaction
        eski_narx = mahsulot.narx or Decimal('0')
        eski_sotib = mahsulot.sotib_olish_narxi or Decimal('0')
        try:
            yangi_narx = _decimal(request.POST.get('narx', eski_narx))
            yangi_sotib = _decimal(request.POST.get('sotib_olish_narxi', eski_sotib))
        except ValueError as e:
            messages.error(request, str(e))
            return redirect('mahsulot_tahrir', mahsulot_id=mahsulot.id)
        if yangi_narx < 0 or yangi_sotib < 0:
            messages.error(request, "Narx manfiy bo'lishi mumkin emas.")
            return redirect('mahsulot_tahrir', mahsulot_id=mahsulot.id)

        nomi = (request.POST.get('nomi') or mahsulot.nomi).strip() or mahsulot.nomi
        birlik = (request.POST.get('birlik') or mahsulot.birlik).strip() or 'dona'
        try:
            minimum_son = _miqdor(request.POST.get('minimum_son', mahsulot.minimum_son), birlik)
        except ValueError:
            minimum_son = mahsulot.minimum_son
        faol = request.POST.get('faol') == 'on'

        with transaction.atomic():
            mahsulot.nomi = nomi
            mahsulot.narx = yangi_narx
            mahsulot.sotib_olish_narxi = yangi_sotib
            mahsulot.birlik = birlik
            mahsulot.minimum_son = minimum_son
            mahsulot.faol = faol
            mahsulot.save(update_fields=['nomi', 'narx', 'sotib_olish_narxi',
                                         'birlik', 'minimum_son', 'faol', 'yangilangan'])
            # Narx yoki tannarx o'zgargan bo'lsa — tarixga yozamiz
            if yangi_narx != eski_narx or yangi_sotib != eski_sotib:
                NarxTarix.objects.create(
                    mahsulot=mahsulot, kim=request.user,
                    eski_narx=eski_narx, yangi_narx=yangi_narx,
                    eski_sotib=eski_sotib, yangi_sotib=yangi_sotib,
                )
        messages.success(request, "✅ Mahsulot saqlandi.")
        return redirect('mahsulot_tahrir', mahsulot_id=mahsulot.id)

    tarix = mahsulot.narx_tarixi.select_related('kim')[:20]
    return render(request, 'mahsulot/tahrir.html', {
        'mahsulot': mahsulot,
        'tarix': tarix,
    })


# ══════════════════════════════════════════════════════════════
#  INVENTARIZATSIYA (ombor reviziyasi / sanoq)
# ══════════════════════════════════════════════════════════════
@login_required
def inventarizatsiya_royxat(request):
    """Sanoq sessiyalari ro'yxati + yangi sanoq boshlash. FAQAT egasi."""
    dokon = _egasi_dokon(request.user)
    if not dokon:
        messages.error(request, "Inventarizatsiya faqat do'kon egasiga ruxsat etilgan.")
        return redirect('mahsulot_royxat')
    ochiq = Inventarizatsiya.objects.filter(dokon=dokon, holat='ochiq').first()
    sessiyalar = (Inventarizatsiya.objects
                  .filter(dokon=dokon)
                  .exclude(holat='ochiq')
                  .prefetch_related('itemlar')[:30])
    return render(request, 'mahsulot/inventarizatsiya_royxat.html', {
        'ochiq': ochiq,
        'sessiyalar': sessiyalar,
    })


@login_required
@require_POST
def inventarizatsiya_boshlash(request):
    """Yangi sanoq ochadi: barcha faol mahsulotlardan snapshot (worksheet). FAQAT egasi."""
    from django.db import transaction
    dokon = _egasi_dokon(request.user)
    if not dokon:
        messages.error(request, "Inventarizatsiya faqat do'kon egasiga ruxsat etilgan.")
        return redirect('mahsulot_royxat')

    mavjud = Inventarizatsiya.objects.filter(dokon=dokon, holat='ochiq').first()
    if mavjud:
        messages.info(request, "Sizda ochiq sanoq bor — davom eting.")
        return redirect('inventarizatsiya_detail', inv_id=mavjud.id)

    mahsulotlar = list(Mahsulot.objects.filter(dokon=dokon, faol=True).order_by('nomi'))
    if not mahsulotlar:
        messages.error(request, "Omborda mahsulot yo'q. Avval mahsulot qo'shing.")
        return redirect('mahsulot_royxat')

    with transaction.atomic():
        inv = Inventarizatsiya.objects.create(
            dokon=dokon, boshlagan=request.user,
            izoh=(request.POST.get('izoh') or '').strip(),
        )
        InventarizatsiyaItem.objects.bulk_create([
            InventarizatsiyaItem(
                inventarizatsiya=inv, mahsulot=m, mahsulot_nomi=m.nomi,
                birlik=m.birlik, tizim_soni=m.soni,
                sotib_olish_narxi=m.sotib_olish_narxi or 0,
            ) for m in mahsulotlar
        ])
    messages.success(request, f"✅ Sanoq boshlandi — {len(mahsulotlar)} ta mahsulot.")
    return redirect('inventarizatsiya_detail', inv_id=inv.id)


@login_required
def inventarizatsiya_detail(request, inv_id):
    """Sanoq varaqasi: real sonlarni kiritish + natija. FAQAT egasi."""
    dokon = _egasi_dokon(request.user)
    if not dokon:
        messages.error(request, "Inventarizatsiya faqat do'kon egasiga ruxsat etilgan.")
        return redirect('mahsulot_royxat')
    inv = get_object_or_404(Inventarizatsiya, id=inv_id, dokon=dokon)
    itemlar = inv.itemlar.select_related('mahsulot').all()
    return render(request, 'mahsulot/inventarizatsiya_detail.html', {
        'inv': inv,
        'itemlar': itemlar,
        'ochiq': inv.holat == 'ochiq',
    })


@login_required
@require_POST
def inventarizatsiya_saqlash(request, inv_id):
    """Kiritilgan real sonlarni saqlaydi. action=saqla|yakunla|bekor.
    - saqla: faqat qiymatlarni yozadi (davom etish mumkin)
    - yakunla: ombor real songa TUZATILADI + log, sessiya yopiladi
    - bekor: ombor o'zgarmaydi, sessiya bekor qilinadi"""
    from django.db import transaction
    from django.utils import timezone
    dokon = _egasi_dokon(request.user)
    if not dokon:
        messages.error(request, "Inventarizatsiya faqat do'kon egasiga ruxsat etilgan.")
        return redirect('mahsulot_royxat')
    inv = get_object_or_404(Inventarizatsiya, id=inv_id, dokon=dokon)
    if inv.holat != 'ochiq':
        messages.error(request, "Bu sanoq allaqachon yakunlangan.")
        return redirect('inventarizatsiya_detail', inv_id=inv.id)

    action = request.POST.get('action', 'saqla')

    if action == 'bekor':
        inv.holat = 'bekor'
        inv.yakunlangan_vaqt = timezone.now()
        inv.save(update_fields=['holat', 'yakunlangan_vaqt'])
        messages.info(request, "Sanoq bekor qilindi — ombor o'zgarmadi.")
        return redirect('inventarizatsiya_royxat')

    # Kiritilgan real sonlarni o'qib, itemlarga yozamiz
    itemlar = list(inv.itemlar.all())
    with transaction.atomic():
        for it in itemlar:
            raw = request.POST.get(f'real_{it.id}', '')
            raw = (raw or '').strip()
            birlik = it.birlik or (it.mahsulot.birlik if it.mahsulot_id else 'dona')
            if raw == '':
                real = None
            else:
                try:
                    real = _miqdor(raw, birlik)
                except ValueError:
                    real = it.real_soni  # noto'g'ri kiritilsa eskisini saqlaymiz
            izoh = (request.POST.get(f'izoh_{it.id}', '') or '').strip()[:200]
            it.real_soni = real
            it.izoh = izoh
        InventarizatsiyaItem.objects.bulk_update(itemlar, ['real_soni', 'izoh'])

        if action == 'yakunla':
            tuzatildi = 0
            for it in itemlar:
                if it.real_soni is None:
                    continue  # sanalmagan — tegmaymiz
                if not it.mahsulot_id:
                    continue
                farq = it.real_soni - it.tizim_soni
                # Omborni real songa tuzatamiz (mutlaq qiymatga o'rnatamiz)
                Mahsulot.objects.filter(id=it.mahsulot_id).update(soni=it.real_soni)
                if farq != 0:
                    tuzatildi += 1
                    # Ombor harakati logi (kirim = ortiqcha, chiqim = kamomad)
                    SkannerLog.objects.create(
                        mahsulot_id=it.mahsulot_id,
                        harakat='kirim' if farq > 0 else 'chiqim',
                        soni=abs(farq), foydalanuvchi=request.user,
                    )
            inv.holat = 'yakunlangan'
            inv.yakunlangan_vaqt = timezone.now()
            inv.save(update_fields=['holat', 'yakunlangan_vaqt'])
            messages.success(
                request,
                f"✅ Sanoq yakunlandi. {tuzatildi} ta mahsulot ombori tuzatildi.")
            return redirect('inventarizatsiya_detail', inv_id=inv.id)

    messages.success(request, "💾 Saqlandi. Sanoqni keyinroq davom ettirishingiz mumkin.")
    return redirect('inventarizatsiya_detail', inv_id=inv.id)
